"""Excel/CSV bulk upload for vendor spare parts (SparePartsImport equivalent)."""

from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timezone

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.models import Part, PartCategory, PartInventory, PartSupplier, Vendor
from app.part_condition_services import parse_part_condition
from app.tier_services import ensure_products_capacity, require_excel_upload, require_vin_decoder
from app.part_image_services import import_part_image_from_url, is_valid_image_url
from app.parts_services import create_part
from app.vin_compat_services import add_part_vin_compatibilities, parse_vin_prefixes

BULK_COLUMNS = [
    "oem_number",
    "name_ar",
    "name_en",
    "part_brand",
    "price",
    "quantity",
    "sub_category",
    "compatible_vehicles",
    "vin_prefixes",
    "description",
    "image_url",
    "condition",
]

TEMPLATE_CSV = (
    "oem_number,name_ar,name_en,part_brand,price,quantity,"
    "sub_category,compatible_vehicles,vin_prefixes,description,image_url,condition\n"
    "# حالة القطعة (condition): new للجديد | used أو مستعمل أو مستعملة للمستعمل\n"
    "TOY-04152-YZZA1,فلتر زيت تويوتا كامري,Toyota Camry Oil Filter,Toyota OEM,45,10,"
    "maintenance-filters,Toyota Camry 2018-2024,4T1B11HK5JK,فلتر أصلي للكامري,"
    "https://example.com/parts/toyota-filter.jpg,new\n"
    "BOSCH-USED-001,فحمات فرامل مستعملة,Used Brake Pads,Bosch,120,2,"
    "brakes,Toyota Corolla,,,,مستعمل\n"
)

MAX_BULK_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".csv", ".xlsx"}


@dataclass
class RowError:
    row: int
    field: str
    message: str


@dataclass
class BulkUploadResult:
    total_rows: int = 0
    imported: int = 0
    errors: list[RowError] = field(default_factory=list)
    used_uncategorized: int = 0

    def to_dict(self) -> dict:
        return {
            "total_rows": self.total_rows,
            "imported": self.imported,
            "failed": len(self.errors),
            "used_uncategorized": self.used_uncategorized,
            "errors": [
                {"row": e.row, "field": e.field, "message": e.message} for e in self.errors
            ],
        }


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _normalize_header(key: str) -> str:
    return key.strip().lower().replace(" ", "_")


def _slugify(value: str, *, prefix: str = "") -> str:
    s = value.lower().strip()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    s = re.sub(r"-+", "-", s).strip("-")
    if not s:
        s = "item"
    if prefix:
        s = f"{prefix}-{s}"
    if not s[0].isalpha():
        s = f"p-{s}"
    return s[:80]


def _parse_compatible_vehicles(raw: str | None) -> list[dict]:
    if not raw or not str(raw).strip():
        return []
    items = [p.strip() for p in re.split(r"[,;]+", str(raw)) if p.strip()]
    return [{"label": item} for item in items]


def _parse_rows_from_csv(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("الملف فارغ أو بدون عناوين أعمدة")
    rows: list[dict[str, str]] = []
    for row in reader:
        normalized = {
            _normalize_header(k): (v or "").strip()
            for k, v in row.items()
            if k
        }
        if normalized.get("oem_number", "").startswith("#"):
            continue
        if any(normalized.values()):
            rows.append(normalized)
    return rows


def _parse_rows_from_xlsx(content: bytes) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise ValueError("ملف Excel بدون أوراق عمل")
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("ملف Excel فارغ") from exc
    headers = [_normalize_header(str(h or "")) for h in header_row]
    if not any(headers):
        raise ValueError("ملف Excel بدون عناوين أعمدة")
    rows: list[dict[str, str]] = []
    for row in rows_iter:
        values = [str(c).strip() if c is not None else "" for c in row]
        if not any(values):
            continue
        padded = values + [""] * max(0, len(headers) - len(values))
        normalized = dict(zip(headers, padded[: len(headers)]))
        rows.append(normalized)
    return rows


def parse_upload_file(content: bytes, filename: str) -> list[dict[str, str]]:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("الصيغ المدعومة: .csv و .xlsx فقط")
    if ext == ".csv":
        return _parse_rows_from_csv(content)
    return _parse_rows_from_xlsx(content)


class _CategoryResolver:
    def __init__(self, db: Session) -> None:
        cats = db.scalars(
            select(PartCategory).where(PartCategory.is_active.is_(True))
        ).all()
        self.by_slug: dict[str, uuid.UUID] = {}
        self.by_name_ar: dict[str, uuid.UUID] = {}
        self.by_name_en: dict[str, uuid.UUID] = {}
        self.uncategorized_id: uuid.UUID | None = None
        for cat in cats:
            self.by_slug[cat.slug.lower()] = cat.id
            self.by_name_ar[cat.name_ar.strip()] = cat.id
            if cat.name_en:
                self.by_name_en[cat.name_en.strip().lower()] = cat.id
            if cat.slug == "uncategorized":
                self.uncategorized_id = cat.id
        if not self.uncategorized_id:
            fallback = self.by_slug.get("uncategorized")
            if fallback:
                self.uncategorized_id = fallback

    def resolve(self, raw: str) -> tuple[uuid.UUID, bool]:
        key = raw.strip()
        if not key:
            if not self.uncategorized_id:
                raise ValueError("قسم غير مصنف غير متوفر في النظام")
            return self.uncategorized_id, True
        lowered = key.lower()
        if lowered in self.by_slug:
            return self.by_slug[lowered], False
        if key in self.by_name_ar:
            return self.by_name_ar[key], False
        if lowered in self.by_name_en:
            return self.by_name_en[lowered], False
        if not self.uncategorized_id:
            raise ValueError("قسم غير مصنف غير متوفر في النظام")
        return self.uncategorized_id, True


def _resolve_supplier(db: Session, brand: str) -> uuid.UUID | None:
    brand = brand.strip()
    if not brand:
        return None
    slug = _slugify(brand, prefix="brand")[:40]
    existing = db.scalar(
        select(PartSupplier).where(
            or_(
                func.lower(PartSupplier.name_ar) == brand.lower(),
                func.lower(PartSupplier.name_en) == brand.lower(),
                PartSupplier.slug == slug,
            )
        )
    )
    if existing:
        return existing.id
    supplier = PartSupplier(
        id=uuid.uuid4(),
        slug=slug,
        name_ar=brand,
        name_en=brand,
        is_oem=False,
        is_active=True,
        created_at=_now(),
    )
    db.add(supplier)
    db.flush()
    return supplier.id


def _validate_row(
    row: dict[str, str],
    row_num: int,
    *,
    seen_oem: set[str],
    seen_slugs: set[str],
    existing_parts: set[str],
    existing_slugs: set[str],
    category_resolver: _CategoryResolver,
) -> tuple[dict | None, list[RowError], bool]:
    errors: list[RowError] = []
    oem = row.get("oem_number", "").strip()
    price_raw = row.get("price", "").strip()
    qty_raw = row.get("quantity", "").strip()
    name_ar = row.get("name_ar", "").strip()

    if not oem:
        errors.append(RowError(row_num, "oem_number", "رقم OEM مطلوب"))
    if not price_raw:
        errors.append(RowError(row_num, "price", "السعر مطلوب"))
    if not qty_raw:
        errors.append(RowError(row_num, "quantity", "الكمية مطلوبة"))
    if not name_ar:
        errors.append(RowError(row_num, "name_ar", "الاسم بالعربية مطلوب"))

    price: float | None = None
    qty: int | None = None
    if price_raw:
        try:
            price = float(price_raw.replace(",", ""))
            if price <= 0:
                errors.append(RowError(row_num, "price", "السعر يجب أن يكون أكبر من صفر"))
        except ValueError:
            errors.append(RowError(row_num, "price", "صيغة السعر غير صالحة"))
    if qty_raw:
        try:
            qty = int(float(qty_raw))
            if qty < 1:
                errors.append(RowError(row_num, "quantity", "الكمية يجب أن تكون 1 على الأقل"))
        except ValueError:
            errors.append(RowError(row_num, "quantity", "صيغة الكمية غير صالحة"))

    if oem:
        oem_key = oem.upper()
        if oem_key in seen_oem:
            errors.append(RowError(row_num, "oem_number", "رقم OEM مكرر في الملف"))
        seen_oem.add(oem_key)
        if oem_key in existing_parts:
            errors.append(
                RowError(row_num, "oem_number", "رقم OEM مسجّل مسبقاً في الكتالوج")
            )
        slug = _slugify(oem)
        if slug in seen_slugs or slug in existing_slugs:
            errors.append(RowError(row_num, "oem_number", "المعرّف (slug) مكرر"))
        seen_slugs.add(slug)

    vin_prefixes: list[str] = []
    vin_raw = row.get("vin_prefixes", "").strip()
    if vin_raw:
        try:
            vin_prefixes = parse_vin_prefixes(vin_raw)
        except ValueError as exc:
            errors.append(RowError(row_num, "vin_prefixes", str(exc)))

    used_uncategorized = False
    category_id: uuid.UUID | None = None
    sub_raw = row.get("sub_category", "").strip()
    try:
        category_id, used_uncategorized = category_resolver.resolve(sub_raw)
    except ValueError as exc:
        errors.append(RowError(row_num, "sub_category", str(exc)))

    image_url = row.get("image_url", "").strip()
    if image_url and not is_valid_image_url(image_url):
        errors.append(RowError(row_num, "image_url", "رابط الصورة غير صالح"))

    if errors:
        return None, errors, False

    assert price is not None and qty is not None and category_id is not None
    parsed = {
        "oem_number": oem,
        "part_number": oem,
        "slug": _slugify(oem),
        "name_ar": name_ar,
        "name_en": row.get("name_en", "").strip() or None,
        "description_ar": row.get("description", "").strip() or None,
        "price_sar": price,
        "qty_available": qty,
        "category_id": category_id,
        "part_brand": row.get("part_brand", "").strip(),
        "compatible_vehicles": _parse_compatible_vehicles(row.get("compatible_vehicles")),
        "vin_prefixes": vin_prefixes,
        "image_url": image_url or None,
        "part_condition": parse_part_condition(row.get("condition")),
        "used_uncategorized": used_uncategorized,
    }
    return parsed, [], used_uncategorized


def process_bulk_upload(
    db: Session,
    vendor: Vendor,
    content: bytes,
    filename: str,
) -> BulkUploadResult:
    if vendor.status != "approved":
        raise ValueError("المحل غير معتمد بعد")
    require_excel_upload(db, vendor)
    if len(content) > MAX_BULK_BYTES:
        raise ValueError("حجم الملف يتجاوز الحد المسموح (10 ميجابايت)")

    raw_rows = parse_upload_file(content, filename)
    result = BulkUploadResult(total_rows=len(raw_rows))
    if not raw_rows:
        raise ValueError("لا توجد صفوف بيانات في الملف")

    ensure_products_capacity(db, vendor, additional=len(raw_rows))
    if any((row.get("vin_prefixes") or "").strip() for row in raw_rows):
        require_vin_decoder(db, vendor)

    category_resolver = _CategoryResolver(db)
    existing_parts = {
        p.upper()
        for p in db.scalars(select(Part.part_number)).all()
    }
    seen_oem: set[str] = set()
    seen_slugs: set[str] = set()
    existing_slugs = {s for s in db.scalars(select(Part.slug)).all()}
    validated: list[tuple[int, dict]] = []

    for idx, row in enumerate(raw_rows, start=2):
        parsed, row_errors, used_uncat = _validate_row(
            row,
            idx,
            seen_oem=seen_oem,
            seen_slugs=seen_slugs,
            existing_parts=existing_parts,
            existing_slugs=existing_slugs,
            category_resolver=category_resolver,
        )
        result.errors.extend(row_errors)
        if parsed:
            validated.append((idx, parsed))
            if used_uncat:
                result.used_uncategorized += 1

    if result.errors:
        return result

    try:
        for _row_num, data in validated:
            supplier_id = _resolve_supplier(db, data["part_brand"])
            prefixes = data["vin_prefixes"]
            part = create_part(
                db,
                part_number=data["part_number"],
                slug=data["slug"],
                name_ar=data["name_ar"],
                name_en=data["name_en"],
                description_ar=data["description_ar"],
                category_id=data["category_id"],
                price_sar=data["price_sar"],
                supplier_id=supplier_id,
                oem_number=data["oem_number"],
                vin_prefix=prefixes[0] if prefixes else None,
                vehicle_compatibility=data["compatible_vehicles"],
                part_condition=data["part_condition"],
                require_leaf_category=True,
            )
            if prefixes:
                add_part_vin_compatibilities(db, part.id, prefixes)
            if data.get("image_url"):
                import_part_image_from_url(db, part, data["image_url"])
            inv = PartInventory(
                id=uuid.uuid4(),
                vendor_id=vendor.id,
                part_id=part.id,
                qty_available=data["qty_available"],
                cost_sar=round(data["price_sar"] * 0.75, 2),
                updated_at=_now(),
            )
            db.add(inv)
            result.imported += 1
        db.commit()
    except Exception:
        db.rollback()
        raise

    return result
